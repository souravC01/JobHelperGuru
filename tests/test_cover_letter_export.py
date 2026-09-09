import io
import docx
import pytest
from fastapi.testclient import TestClient
from backend.main import app
from backend.models import User
from backend.routers.auth import get_current_user


def test_cover_letter_docx_export_unauthenticated():
    client = TestClient(app)
    resp = client.post(
        "/api/resumes/export-cover-letter-docx",
        json={
            "cover_letter_text": "Dear Team,\n\nParagraph 1\n\nParagraph 2\n\nParagraph 3\n\nSincerely,\nAlice",
            "company": "Acme",
            "role": "Staff Engineer",
        },
    )
    assert resp.status_code == 401


def test_cover_letter_docx_export_validation_error():
    client = TestClient(app)
    app.dependency_overrides[get_current_user] = lambda: User(
        id="user-docx-1", email="docx@example.com", name="Alice Developer", created_at="2026-09-07T00:00:00"
    )
    try:
        resp = client.post(
            "/api/resumes/export-cover-letter-docx",
            json={
                "cover_letter_text": "   ",
                "company": "Acme",
            },
        )
        assert resp.status_code == 422
    finally:
        app.dependency_overrides.pop(get_current_user, None)


def test_cover_letter_docx_export_success_and_parses():
    client = TestClient(app)
    app.dependency_overrides[get_current_user] = lambda: User(
        id="user-docx-1", email="docx@example.com", name="Alice Developer", created_at="2026-09-07T00:00:00"
    )
    try:
        payload = {
            "cover_letter_text": (
                "Dear Hiring Team at Stripe,\n\n"
                "I am writing to express my enthusiasm for the Backend Engineer position at Stripe.\n\n"
                "My experience designing distributed event architectures and low-latency APIs aligns directly with Stripe's payments infrastructure.\n\n"
                "I welcome the opportunity to discuss how my engineering background will accelerate your roadmap.\n\n"
                "Sincerely,\nAlice Developer"
            ),
            "company": "Stripe",
            "role": "Backend Engineer",
            "candidate_name": "Alice Developer",
            "subject_line": "Application: Backend Engineer - Alice Developer",
        }
        resp = client.post("/api/resumes/export-cover-letter-docx", json=payload)
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in resp.headers["content-type"]
        assert 'attachment; filename="Cover_Letter_Stripe_Backend_Engineer.docx"' in resp.headers["content-disposition"]

        # Parse with python-docx
        doc = docx.Document(io.BytesIO(resp.content))
        doc_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        assert "Alice Developer" in doc_text
        assert "Stripe" in doc_text
        assert "Backend Engineer" in doc_text
        assert "distributed event architectures" in doc_text
        assert "\u2014" not in doc_text
        assert "\u2013" not in doc_text
    finally:
        app.dependency_overrides.pop(get_current_user, None)


def test_resume_download_fallback_when_binary_file_missing():
    client = TestClient(app)
    app.dependency_overrides[get_current_user] = lambda: User(
        id="user-fb-1", email="fallback@example.com", name="Fallback User", created_at="2026-09-07T00:00:00"
    )
    try:
        from backend.main import storage
        # Create a resume that has raw text but no physical file attached
        res = storage.add_resume(
            name="Text-Only Resume",
            content="Senior Software Engineer with 8 years of Python experience.\nLed distributed systems team.",
            user_id="user-fb-1",
        )
        resume_id = res.id

        resp = client.get(f"/api/resumes/{resume_id}/download")
        assert resp.status_code == 200
        assert resp.headers.get("x-fallback-generated") == "true"
        assert "Text-Only_Resume_generated.docx" in resp.headers.get("content-disposition", "")

        # Verify generated document is readable docx
        doc = docx.Document(io.BytesIO(resp.content))
        doc_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        assert "Text-Only Resume" in doc_text
        assert "Senior Software Engineer" in doc_text
    finally:
        app.dependency_overrides.pop(get_current_user, None)
