import io
from itertools import pairwise

import docx
import pypdf
import pytest


def test_chunking_preserves_tail_and_original_offsets():
    from backend.services.matching.chunking import chunk_text

    text = "Intro\n" + "x" * 12000 + "\nBuilt production Kafka consumers."
    chunks = chunk_text(text)

    assert chunks[0].start == 0 and chunks[-1].end == len(text)
    assert all(chunk.text == text[chunk.start:chunk.end] for chunk in chunks)
    assert all(next_chunk.start <= chunk.end for chunk, next_chunk in pairwise(chunks))
    assert "Kafka" in chunks[-1].text


def test_chunking_covers_every_character_and_rejects_oversized_input():
    from backend.services.matching.chunking import MAX_DOCUMENT_CHARS, chunk_text

    text = "First paragraph.\n\n" + "second paragraph " * 1000 + "\n\nFinal paragraph."
    chunks = chunk_text(text)
    covered = [False] * len(text)
    for chunk in chunks:
        for index in range(chunk.start, chunk.end):
            covered[index] = True

    assert all(covered)
    assert all(len(chunk.text) <= 8000 for chunk in chunks)
    with pytest.raises(ValueError, match="100,000"):
        chunk_text("x" * (MAX_DOCUMENT_CHARS + 1))


@pytest.mark.parametrize("delimiter", ["\n", "\n\n"])
def test_chunking_never_extends_past_the_maximum_for_a_boundary_at_the_limit(delimiter):
    from backend.services.matching.chunking import MAX_CHUNK_CHARS, chunk_text

    chunks = chunk_text("x" * MAX_CHUNK_CHARS + delimiter + "tail")

    assert all(len(chunk.text) <= MAX_CHUNK_CHARS for chunk in chunks)


def test_pdf_final_text_limit_counts_page_separators(monkeypatch):
    from backend.services.document_parser import (
        MAX_EXTRACTED_CHARS,
        extract_text_with_warnings,
    )

    class Page:
        def __init__(self, text):
            self.text = text

        def extract_text(self):
            return self.text

    class Reader:
        def __init__(self, _stream):
            self.pages = [Page("x" * 50_000), Page("y" * 49_998)]

    monkeypatch.setattr(pypdf, "PdfReader", Reader)
    exact = extract_text_with_warnings(b"pdf", "resume.pdf")

    assert len(exact.text) == MAX_EXTRACTED_CHARS

    class OverflowReader:
        def __init__(self, _stream):
            self.pages = [Page("x" * 50_000), Page("y" * 49_999)]

    monkeypatch.setattr(pypdf, "PdfReader", OverflowReader)
    with pytest.raises(ValueError, match="100,000"):
        extract_text_with_warnings(b"pdf", "resume.pdf")


def test_docx_final_text_limit_counts_paragraph_separators():
    from backend.services.document_parser import (
        MAX_EXTRACTED_CHARS,
        extract_text_with_warnings,
    )

    def document_bytes(second_paragraph):
        document = docx.Document()
        document.add_paragraph("x" * 50_000)
        document.add_paragraph(second_paragraph)
        stream = io.BytesIO()
        document.save(stream)
        return stream.getvalue()

    exact = extract_text_with_warnings(document_bytes("y" * 49_998), "resume.docx")

    assert len(exact.text) == MAX_EXTRACTED_CHARS
    with pytest.raises(ValueError, match="100,000"):
        extract_text_with_warnings(document_bytes("y" * 49_999), "resume.docx")


def test_document_parser_exposes_pdf_extraction_warnings_and_keeps_string_wrapper():
    from backend.services.document_parser import (
        extract_text_from_file,
        extract_text_with_warnings,
    )

    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    stream = io.BytesIO()
    writer.write(stream)

    result = extract_text_with_warnings(stream.getvalue(), "blank.pdf")

    assert result.text == ""
    assert result.warnings == ["PDF page 1 did not contain extractable text."]
    assert extract_text_from_file(stream.getvalue(), "blank.pdf") == result.text


def test_parse_endpoint_returns_non_utf8_decoding_warnings(client):
    registration = client.post(
        "/api/auth/register",
        json={
            "email": "parser-warnings@example.test",
            "password": "Password123!",
            "name": "Parser Warnings",
        },
    )
    headers = {"Authorization": f"Bearer {registration.json()['token']}"}

    response = client.post(
        "/api/resumes/parse-file",
        headers=headers,
        files={"file": ("resume.txt", io.BytesIO(b"\x96Python"), "text/plain")},
    )

    assert response.status_code == 200
    assert response.json()["text"] == "\x96Python"
    assert response.json()["warnings"] == [
        "Document text was decoded as Latin-1 after UTF-8 decoding failed."
    ]
