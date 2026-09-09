import io
import re
import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from backend.services.ai_engine import sanitize_dashes


def generate_cover_letter_docx(
    cover_letter_text: str,
    candidate_name: str = "",
    company: str = "",
    role: str = "",
    subject_line: str = "",
) -> bytes:
    """
    Generates an executive-formatted Word document (.docx) for the 3-paragraph cover letter.
    Enforces zero em-dashes and standard margins.
    """
    doc = docx.Document()

    # Page Margins: 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal Style
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    clean_text = sanitize_dashes(cover_letter_text or "").strip()
    clean_name = sanitize_dashes(candidate_name or "").strip()
    clean_company = sanitize_dashes(company or "").strip()
    clean_role = sanitize_dashes(role or "").strip()
    clean_subject = sanitize_dashes(subject_line or "").strip()

    # Candidate Name Header (if available)
    if clean_name:
        name_p = doc.add_paragraph()
        name_p.paragraph_format.space_before = Pt(0)
        name_p.paragraph_format.space_after = Pt(4)
        run = name_p.add_run(clean_name)
        run.bold = True
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(15, 23, 42)

    # Date
    today_str = datetime.date.today().strftime("%B %d, %Y")
    date_p = doc.add_paragraph()
    date_p.paragraph_format.space_before = Pt(0)
    date_p.paragraph_format.space_after = Pt(14)
    run_date = date_p.add_run(today_str)
    run_date.font.size = Pt(10)
    run_date.font.color.rgb = RGBColor(100, 116, 139)

    # Recipient block
    if clean_company or clean_role:
        recip_p = doc.add_paragraph()
        recip_p.paragraph_format.space_before = Pt(0)
        recip_p.paragraph_format.space_after = Pt(12)
        recip_p.paragraph_format.line_spacing = 1.15
        recip_p.add_run("Hiring Team\n")
        if clean_company:
            recip_p.add_run(f"{clean_company}\n")

    # Subject line
    if clean_subject:
        subj_p = doc.add_paragraph()
        subj_p.paragraph_format.space_before = Pt(4)
        subj_p.paragraph_format.space_after = Pt(12)
        subj_run = subj_p.add_run(f"Subject: {clean_subject}")
        subj_run.bold = True
        subj_run.font.color.rgb = RGBColor(15, 23, 42)

    # Body paragraphs
    paragraphs = [p.strip() for p in clean_text.split("\n\n") if p.strip()]
    for p_content in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.15
        p.add_run(p_content)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def generate_docx_from_text(title: str, text: str) -> bytes:
    """
    Generates a cleanly formatted Word document (.docx) from plain text.
    Used for graceful fallback when the original binary file is missing from storage.
    """
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    clean_title = sanitize_dashes(title or "Document").strip()
    clean_text = sanitize_dashes(text or "").strip()

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(12)
    run = title_p.add_run(clean_title)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(15, 23, 42)

    paragraphs = [p.strip() for p in clean_text.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [p.strip() for p in clean_text.split("\n") if p.strip()]

    for p_content in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        p.add_run(p_content)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
